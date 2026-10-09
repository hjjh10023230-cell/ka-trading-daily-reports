#!/usr/bin/env python3
"""Validate and aggregate a local KA trade workbook without copying it into the repo."""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

DETAILED_REQUIRED = ["股票名称", "买卖方向", "ka类型", "用户数", "交易订单数", "trade_amt_usd"]
PIVOT_REQUIRED_DIMENSIONS = ["分区日期", "买卖方向", "证券类型V2", "ka类型"]
PIVOT_OPTIONAL_DIMENSIONS = ["股票名称"]
PIVOT_MEASURES = ["用户数", "交易订单数", "trade_amt_usd"]
DIRECTIONS = {"BUY", "SELL", "BUY_BACK", "SELL_SHORT"}
MARKETS = {"JP", "US"}
OPTION_RE = re.compile(r"(?:\d{6}[CP]\d+|\s\d{6}\s+\d+(?:\.\d+)?[CP])$", re.I)


def number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    return float(value)


def add(bucket: dict[str, float | int], users: float, orders: float, amount: float) -> None:
    bucket["userOccurrences"] += int(users)
    bucket["orders"] += int(orders)
    bucket["amount"] += amount


def empty_bucket() -> dict[str, float | int]:
    return {"userOccurrences": 0, "orders": 0, "amount": 0.0}


def finish_bucket(bucket: dict[str, float | int]) -> dict[str, float | int]:
    orders = int(bucket["orders"])
    return {
        **bucket,
        "averageOrderAmount": float(bucket["amount"]) / orders if orders else 0.0,
    }


def asset_class(code: str, name: str) -> str:
    text = f"{code} {name}".strip()
    if OPTION_RE.search(text):
        return "OPTION"
    if "ETF" in text.upper():
        return "FUND"
    return "EQUITY"


def merged_lookup(ws: Any) -> dict[tuple[int, int], Any]:
    values: dict[tuple[int, int], Any] = {}
    for area in ws.merged_cells.ranges:
        top = ws.cell(area.min_row, area.min_col).value
        for row in range(area.min_row, area.max_row + 1):
            for column in range(area.min_col, area.max_col + 1):
                values[(row, column)] = top
    return values


def aggregate_rows(rows: list[dict[str, Any]], source_mode: str) -> dict[str, Any]:
    overall = empty_bucket()
    directions = defaultdict(empty_bucket)
    crowds = defaultdict(empty_bucket)
    crowd_directions = defaultdict(lambda: defaultdict(empty_bucket))
    crowd_markets = defaultdict(lambda: defaultdict(empty_bucket))
    crowd_assets = defaultdict(lambda: defaultdict(empty_bucket))
    markets = defaultdict(empty_bucket)
    assets = defaultdict(empty_bucket)
    symbols = defaultdict(lambda: defaultdict(empty_bucket))

    for row in rows:
        users = row["users"]
        orders = row["orders"]
        amount = row["amount"]
        crowd = row["crowd"]
        direction = row["direction"]
        market = row.get("market", "UNKNOWN")
        klass = row["assetClass"]
        add(overall, users, orders, amount)
        add(directions[direction], users, orders, amount)
        add(crowds[crowd], users, orders, amount)
        add(crowd_directions[crowd][direction], users, orders, amount)
        add(crowd_markets[crowd][market], users, orders, amount)
        add(crowd_assets[crowd][klass], users, orders, amount)
        add(markets[market], users, orders, amount)
        add(assets[klass], users, orders, amount)
        if row.get("symbol"):
            add(symbols[row["symbol"]][direction], users, orders, amount)

    return {
        "sourceMode": source_mode,
        "overallOccurrences": finish_bucket(overall),
        "directions": {key: finish_bucket(value) for key, value in directions.items()},
        "crowds": {
            crowd: {
                **finish_bucket(total),
                "directions": {key: finish_bucket(value) for key, value in crowd_directions[crowd].items()},
                "markets": {key: finish_bucket(value) for key, value in crowd_markets[crowd].items()},
                "assetClasses": {key: finish_bucket(value) for key, value in crowd_assets[crowd].items()},
            }
            for crowd, total in crowds.items()
        },
        "markets": {key: finish_bucket(value) for key, value in markets.items()},
        "assetClasses": {key: finish_bucket(value) for key, value in assets.items()},
        "symbols": {
            symbol: {key: finish_bucket(value) for key, value in values.items()}
            for symbol, values in symbols.items()
        },
    }


def aggregate_detailed(ws: Any) -> dict[str, Any]:
    headers = [cell.value for cell in ws[1]]
    missing = [name for name in DETAILED_REQUIRED if name not in headers]
    if missing:
        raise ValueError(f"缺少字段：{', '.join(missing)}")
    col = {name: headers.index(name) for name in DETAILED_REQUIRED}
    merged = merged_lookup(ws)
    rows: list[dict[str, Any]] = []

    code_column = headers.index("股票代码") if "股票代码" in headers else None
    for row_no, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        values = list(row)
        code = values[code_column] if code_column is not None else ""
        name = values[col["股票名称"]]
        if code_column is not None and code in (None, ""):
            code = merged.get((row_no, code_column + 1), "")
        if name in (None, ""):
            name = merged.get((row_no, col["股票名称"] + 1), "")
        direction = str(values[col["买卖方向"]] or "").upper()
        if direction not in DIRECTIONS:
            raise ValueError(f"第 {row_no} 行存在未知买卖方向：{direction!r}")
        rows.append(
            {
                "crowd": str(values[col["ka类型"]] or "未分类"),
                "direction": direction,
                "market": "UNKNOWN",
                "assetClass": asset_class(str(code), str(name)),
                "symbol": str(code or name),
                "users": number(values[col["用户数"]]),
                "orders": number(values[col["交易订单数"]]),
                "amount": number(values[col["trade_amt_usd"]]),
            }
        )
    return aggregate_rows(rows, "symbol-detail")


def aggregate_pivot(ws: Any, report_date: str) -> dict[str, Any]:
    merged = merged_lookup(ws)
    second_row = [ws.cell(2, column).value for column in range(1, ws.max_column + 1)]
    dimension_columns = {
        name: second_row.index(name) + 1
        for name in PIVOT_REQUIRED_DIMENSIONS + PIVOT_OPTIONAL_DIMENSIONS
        if name in second_row
    }
    missing = [name for name in PIVOT_REQUIRED_DIMENSIONS if name not in dimension_columns]
    if missing:
        raise ValueError(f"双层表头缺少维度：{', '.join(missing)}")
    has_symbol = "股票名称" in dimension_columns

    measure_columns: list[tuple[str, str, int]] = []
    for column in range(1, ws.max_column + 1):
        measure = ws.cell(1, column).value or merged.get((1, column))
        header = ws.cell(2, column).value
        market = str(header or "").upper()
        if measure is None and header in dimension_columns:
            continue
        if measure not in PIVOT_MEASURES or market not in MARKETS:
            raise ValueError(f"第 {column} 列存在未知指标/市场：{measure!r}/{market!r}")
        measure_columns.append((str(measure), market, column))

    rows: list[dict[str, Any]] = []
    seen_dates: set[str] = set()
    for row_no in range(3, ws.max_row + 1):
        def dimension(name: str, default: str = "") -> Any:
            column = dimension_columns[name]
            return ws.cell(row_no, column).value or merged.get((row_no, column)) or default

        raw_date = dimension("分区日期")
        if isinstance(raw_date, (datetime, date)):
            day = raw_date.strftime("%Y-%m-%d")
        else:
            day = str(raw_date or "")
        seen_dates.add(day)
        direction = str(dimension("买卖方向")).upper()
        klass = str(dimension("证券类型V2")).upper()
        crowd = str(dimension("ka类型", "未分类"))
        symbol = str(dimension("股票名称")) if has_symbol else ""
        if direction not in DIRECTIONS:
            raise ValueError(f"第 {row_no} 行存在未知买卖方向：{direction!r}")
        if not klass:
            raise ValueError(f"第 {row_no} 行缺少证券类型")
        by_market = defaultdict(lambda: {name: 0.0 for name in PIVOT_MEASURES})
        for measure, market, column in measure_columns:
            by_market[market][measure] += number(ws.cell(row_no, column).value)
        for market, values in by_market.items():
            if any(values.values()):
                rows.append(
                    {
                        "crowd": crowd,
                        "direction": direction,
                        "market": market,
                        "assetClass": klass,
                        "symbol": symbol,
                        "users": values["用户数"],
                        "orders": values["交易订单数"],
                        "amount": values["trade_amt_usd"],
                    }
                )
    if seen_dates != {report_date}:
        raise ValueError(f"源表日期 {sorted(seen_dates)!r} 与 --date {report_date!r} 不一致")
    return aggregate_rows(rows, "crowd-market-asset-symbol-pivot" if has_symbol else "crowd-market-asset-pivot")


def aggregate(path: Path, report_date: str) -> dict[str, Any]:
    wb = load_workbook(path, read_only=False, data_only=True)
    if "报表" not in wb.sheetnames:
        raise ValueError("缺少工作表：报表")
    ws = wb["报表"]
    first_row = [cell.value for cell in ws[1]]
    second_row = [cell.value for cell in ws[2]] if ws.max_row >= 2 else []
    if all(name in first_row for name in DETAILED_REQUIRED):
        return aggregate_detailed(ws)
    if all(name in second_row for name in PIVOT_REQUIRED_DIMENSIONS):
        return aggregate_pivot(ws, report_date)
    raise ValueError("无法识别工作簿结构：既不是逐标的明细，也不是人群×市场×品类双层透视表")


def load_trend(path: Path, report_date: str) -> dict[str, Any]:
    trend = json.loads(path.read_text(encoding="utf-8"))
    if trend.get("reportDate") != report_date:
        raise ValueError("趋势快照 reportDate 与 --date 不一致")
    if not isinstance(trend.get("crowds"), dict) or not trend["crowds"]:
        raise ValueError("趋势快照缺少 crowds")
    return trend


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path, help="仓库外的本地 Excel 路径")
    parser.add_argument("--date", required=True, help="报告交易日 YYYY-MM-DD")
    parser.add_argument("--trend", type=Path, help="可选：仓库外的分人群滚动日均趋势 JSON")
    parser.add_argument("--output", type=Path, help="可选：将聚合 JSON 写入本地路径")
    args = parser.parse_args()
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", args.date):
        raise SystemExit("--date 必须为 YYYY-MM-DD")
    result = {
        "reportDate": args.date,
        "warning": "方向性成交金额差不等于真实净入金；用户数跨方向、市场和品类不可直接相加去重。",
        **aggregate(args.workbook, args.date),
    }
    if args.trend:
        result["trend"] = load_trend(args.trend, args.date)
    payload = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        args.output.write_text(payload, encoding="utf-8")
    else:
        if hasattr(sys.stdout, "reconfigure"):
            sys.stdout.reconfigure(encoding="utf-8")
        print(payload)


if __name__ == "__main__":
    main()
